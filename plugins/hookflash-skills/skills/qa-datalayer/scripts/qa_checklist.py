"""The dataLayer QA checklist workbook, in the Hookflash house style.

Two tabs. 'QA Checklist' holds one block per event: the event, its category
badge and its status on top, the trigger and a note under that, then one row
per parameter with Present / fired? and Value correct? dropdowns. 'Summary'
lists the blocks, and its counts and statuses roll up live from them. The
layout, palette and formulas are the QA team's own template (their
datalayer-qa-checklist skill); this module is the one builder for it.

The same file ships in two places, and the two copies are kept identical:
Tapa's DataLayer Doc Creator exports the checklist prefilled with the events a
person documented (checks blank, for a tester to fill in), and the qa-datalayer
skill fills the checks in from the pushes it captured. A change to the house
style is a change to both.

Formula cells carry their results. openpyxl never stores them, and without
them Excel's Protected View (which every downloaded file opens in) and the
Outlook / Teams previewers show every status and count blank. Every answer is
known at build time, so save() writes it in; Excel recalculates as soon as
someone edits a check.

Spec (a dict):

    {"client": "Example Store", "batch": "Batch 1", "source": "Example Store DataLayer Guide",
     "events": [{"category": "Non-Ecom", "event": "quiz_progress",
                 "variant": "", "trigger": "...", "note": "...",
                 "params": [{"name": "event", "type": "String",
                             "required": "Yes", "expected": "quiz_progress",
                             "present": "Pass", "value_ok": "Fail",
                             "notes": "..."}]}]}

A param may also be a list, [name, type, required, expected, notes]. present
and value_ok are Pass, Fail, N/A or blank (untested). Optional spec keys:
"instructions" replaces the QA tab's how-to line, "steps" adds lines to the top
of the Summary's How to use.
"""

import io
import math
import re
import zipfile
from xml.sax.saxutils import escape as _xml_escape

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ------------------------------------------------------------------ house style
F = "Arial"
NAVY, SLATE, CYAN, PALE, WHITE, LINE = "1B2A41", "3D5060", "17BEE0", "EEF3F6", "FFFFFF", "C9D2D9"
STATUS_COLOURS = {            # value: (fill, font colour)
    "Pass": ("C6EFCE", "1E6B34"),
    "Fail": ("FFC7CE", "9C0006"),
    "In progress": ("FFEB9C", "9C6500"),
    "N/A": ("D9D9D9", "666666"),
}
CATEGORIES = ("Ecom", "Non-Ecom", "Recurring")
CHECKS = ("Pass", "Fail", "N/A")
REQUIRED = ("Yes", "No", "If available")
PARAM_HEADERS = ["Parameter", "Type", "Required", "Expected value / example",
                 "Present / fired?", "Value correct?", "Notes"]
CHECK_WIDTHS = {"A": 26, "B": 9, "C": 13, "D": 46, "E": 16, "F": 15, "G": 40}
SUMMARY_WIDTHS = {"A": 5, "B": 13, "C": 20, "D": 26, "E": 52, "F": 9, "G": 15}
SUMMARY_SHEET, CHECK_SHEET = "Summary", "QA Checklist"
NO_VARIANT = "-"
DEFAULT_INSTRUCTIONS = ("Fill the Present / fired? and Value correct? columns per parameter. "
                        "Event status rolls up automatically.")
HOW_TO_USE = [
    "1.  Go to the 'QA Checklist' tab. Each push is a block: event details on top, one row per parameter below.",
    "2.  For every parameter set 'Present / fired?' (did the key appear in the dataLayer push?) and "
    "'Value correct?' (is the value right & correctly typed?). Use the dropdown: Pass / Fail / N/A.",
    "3.  Leave a check blank until tested: the Event status shows 'In progress' while any check is blank.",
    "4.  Event status auto-calculates: Fail if any check fails · In progress if any check is untested · "
    "Pass when all checks are Pass or N/A.",
    "5.  Record specifics (actual value seen, screenshot ref, GTM preview note) in the Notes column.",
]

S = Side(style="thin", color=LINE)
BOX = Border(left=S, right=S, top=S, bottom=S)


def fill(c):
    return PatternFill("solid", fgColor=c)


# ------------------------------------------------------------------ text
def put(ws, row, col, value):
    """Write a literal. Text loses the control characters openpyxl refuses,
    and a leading "=" stays text (an expected value, a trigger), never a
    formula."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    c = ws.cell(row, col)
    c.value = value
    if isinstance(value, str) and value.startswith("="):
        c.data_type = "s"
    return c


def lines_for(value, width, font_pt):
    """How many lines Excel wraps `value` onto in a column `width` wide.
    Rendered through Excel, Arial prose fits about width * 12.5 / pt
    characters a line; word wrap leaves lines ragged, so plan on 11."""
    if not value:
        return 1
    per_line = max(1, int(width * 11 / font_pt))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(value).split("\n"))


def height_for(n_lines, font_pt, minimum):
    """Points for `n_lines` of `font_pt` text, never under `minimum` (the
    template's own height for that row). Two lines of 9pt come to the
    template's 25.5."""
    return max(minimum, round(n_lines * font_pt * 4 / 3 + 1.5, 2))


# ------------------------------------------------------------------ spec
def _canon(value, allowed):
    """`value` in its canonical spelling from `allowed`, or "" when blank.
    Unknown values pass through, so a typo shows up in the sheet rather than
    silently becoming blank."""
    if value is None:
        return ""
    s = str(value).strip()
    if not s:
        return ""
    for a in allowed:
        if s.lower() == a.lower() or (a == "N/A" and s.lower() in ("na", "n.a.")):
            return a
    return s


def norm_param(p):
    if isinstance(p, dict):
        d = dict(p)
    else:
        p = list(p) + [None] * (5 - len(p))
        d = {"name": p[0], "type": p[1], "required": p[2], "expected": p[3], "notes": p[4]}
    return {
        "name": d.get("name") or "",
        "type": d.get("type") or "String",
        "required": _canon(d.get("required"), REQUIRED) or "Yes",
        "expected": "" if d.get("expected") is None else str(d.get("expected")),
        "present": _canon(d.get("present"), CHECKS),
        "value_ok": _canon(d.get("value_ok"), CHECKS),
        "notes": d.get("notes") or None,
    }


def check_spec(spec):
    """The spec's problems, as a list of lines (empty = fine)."""
    problems = []
    events = spec.get("events") or []
    if not events:
        problems.append("spec has no events")
    for i, ev in enumerate(events, 1):
        label = f"event {i} ({ev.get('event')})"
        if ev.get("category") not in CATEGORIES:
            problems.append(f"{label}: category must be one of {list(CATEGORIES)}")
        if not ev.get("event"):
            problems.append(f"{label}: no event name")
        if not ev.get("params"):
            problems.append(f"{label}: no params")
        for p in ev.get("params") or []:
            p = norm_param(p)
            for key in ("present", "value_ok"):
                if p[key] and p[key] not in CHECKS:
                    problems.append(f"{label}: {p['name']} {key} is {p[key]!r}, not one of {list(CHECKS)}")
    return problems


def block_status(params):
    """What the block's status formula computes: Fail if any check fails, In
    progress if any listed parameter has a blank check, Pass otherwise. Rows
    with no parameter name are ignored, like the formula ignores them."""
    listed = [p for p in params if p["name"]]
    if not listed:
        return ""
    checks = [p["present"] for p in listed] + [p["value_ok"] for p in listed]
    if any(c == "Fail" for c in checks):
        return "Fail"
    if any(not c for c in checks):
        return "In progress"
    return "Pass"


def block_status_formula(first, last):
    a, e, f = f"A{first}:A{last}", f"E{first}:E{last}", f"F{first}:F{last}"
    return (f'=IF(COUNTA({a})=0,"",IF(COUNTIF(E{first}:F{last},"Fail")>0,"Fail",'
            f'IF(SUMPRODUCT(({a}<>"")*({e}=""))+SUMPRODUCT(({a}<>"")*({f}=""))>0,'
            f'"In progress","Pass")))')


# ------------------------------------------------------------------ builders
def bar(ws, row, value, size, bg, height, bold=True, colour=WHITE, last_col=7):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = put(ws, row, 1, value)
    c.font = Font(name=F, sz=size, bold=bold, color=colour)
    c.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(1, last_col + 1):
        ws.cell(row, col).fill = fill(bg)
    ws.row_dimensions[row].height = height


def status_rules(ws, rng, size, bold, keys):
    """Colour a status cell by its value. A conditional format's solid fill is
    drawn from its bgColor: with only fgColor set (fill()), Excel shows the
    font colour and no fill."""
    for k in keys:
        bg, fc = STATUS_COLOURS[k]
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=[f'"{k}"'],
                            fill=PatternFill("solid", start_color=bg, end_color=bg),
                            font=Font(name=F, sz=size, bold=bold, color=fc)))


def _heading(spec):
    client = (spec.get("client") or "").strip() or "Client"
    batch = (spec.get("batch") or "").strip()
    source = (spec.get("source") or "").strip() or f"{client} DataLayer Guide"
    return client, batch, source


def build_checklist(wb, spec, events):
    """The 'QA Checklist' tab. Returns one (status_cell, header_row, first,
    last) per block, and the cached result of each status formula."""
    client, batch, source = _heading(spec)
    ws = wb.create_sheet(CHECK_SHEET)
    for k, v in CHECK_WIDTHS.items():
        ws.column_dimensions[k].width = v
    bar(ws, 1, f"{client} · dataLayer QA Checklist", 16, NAVY, 27.75)
    strap = "  ·  ".join(x for x in (batch, f"Source: {source}",
                                     spec.get("instructions") or DEFAULT_INSTRUCTIONS) if x)
    bar(ws, 2, strap, 9, SLATE, height_for(lines_for(strap, sum(CHECK_WIDTHS.values()), 9), 9, 21.75),
        bold=False)
    ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    r, out, cached = 4, [], {}
    for i, ev in enumerate(events, 1):
        params = ev["params"]
        first = r + 3
        last = first + len(params) - 1

        variant = (ev.get("variant") or "").strip()
        title = f"{i}.  {ev['event']}" + (f"   ·   {variant}" if variant and variant != NO_VARIANT else "")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = put(ws, r, 1, title)
        c.font = Font(name=F, sz=12, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="left", vertical="center")
        for col in range(1, 5):
            ws.cell(r, col).fill = fill(NAVY)
        c = put(ws, r, 5, ev["category"])
        c.font = Font(name=F, sz=10, bold=True, color=WHITE)
        c.fill = fill(CYAN)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c = ws.cell(r, 6, block_status_formula(first, last))
        c.font = Font(name=F, sz=10, bold=True, color="000000")
        c.fill = fill(WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
        status_rules(ws, f"F{r}", 10, True, STATUS_COLOURS)
        cached[f"F{r}"] = block_status(params)
        ws.row_dimensions[r].height = 21.75

        strip = f"Trigger: {ev.get('trigger') or ''}"
        if ev.get("note"):
            strip += f"   |   Note: {ev['note']}"
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)
        c = put(ws, r + 1, 1, strip)
        c.font = Font(name=F, sz=9, color="333333")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        for col in range(1, 6):
            ws.cell(r + 1, col).fill = fill(PALE)
        strip_width = sum(CHECK_WIDTHS[k] for k in "ABCDE")
        ws.row_dimensions[r + 1].height = height_for(lines_for(strip, strip_width, 9), 9, 30)

        for col, h in enumerate(PARAM_HEADERS, 1):
            c = ws.cell(r + 2, col, h)
            c.font = Font(name=F, sz=9, bold=True, color=WHITE)
            c.fill = fill(SLATE)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BOX
        ws.row_dimensions[r + 2].height = 25.5

        for j, p in enumerate(params):
            rr = first + j
            band = PALE if j % 2 == 1 else WHITE
            values = [p["name"], p["type"], p["required"], p["expected"],
                      p["present"] or None, p["value_ok"] or None, p["notes"]]
            for col, v in enumerate(values, 1):
                c = put(ws, rr, col, v)
                c.border = BOX
                c.fill = fill(PALE if col in (5, 6) else band)
                if col == 1:
                    c.font = Font(name=F, sz=10, bold=True)
                    c.alignment = Alignment(horizontal="left", vertical="center")
                elif col in (4, 7):
                    c.font = Font(name=F, sz=9)
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                else:
                    c.font = Font(name=F, sz=9)
                    c.alignment = Alignment(horizontal="center", vertical="center")
            # Only a row whose value or note wraps gets a height; a one-line
            # row keeps Excel's default, as in the template.
            wraps = max(lines_for(p["expected"], CHECK_WIDTHS["D"], 9),
                        lines_for(p["notes"], CHECK_WIDTHS["G"], 9))
            if wraps > 1:
                ws.row_dimensions[rr].height = height_for(wraps, 9, 15)
        out.append((f"F{r}", r, first, last))
        r = last + 2

    dv = DataValidation(type="list", formula1='"' + ",".join(CHECKS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for _, _, first, last in out:
        dv.add(f"E{first}:F{last}")
        status_rules(ws, f"E{first}:F{last}", 9, False, list(CHECKS))
    ws.sheet_view.showGridLines = False
    return ws, out, cached


def build_summary(wb, spec, events, blocks, statuses):
    client, batch, _source = _heading(spec)
    sm = wb.create_sheet(SUMMARY_SHEET, 0)
    for k, v in SUMMARY_WIDTHS.items():
        sm.column_dimensions[k].width = v
    bar(sm, 1, f"{client} · dataLayer QA", 18, NAVY, 33.75)
    scope = f"the {batch.lower()} dataLayer pushes" if batch else "the dataLayer pushes"
    bar(sm, 2, f"Quality-assurance checklist for {scope}. "
               "Statuses below roll up live from the 'QA Checklist' tab.", 10, SLATE, 21.75, bold=False)
    c = sm["A4"]
    c.value = "How to use"
    c.font = Font(name=F, sz=12, bold=True, color=NAVY)

    row = 5
    width = sum(SUMMARY_WIDTHS.values())
    for step in list(spec.get("steps") or []) + HOW_TO_USE:
        sm.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = put(sm, row, 1, step)
        c.font = Font(name=F, sz=9)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        sm.row_dimensions[row].height = height_for(lines_for(step, width, 9), 9, 25.5)
        row += 1

    kr = row + 1
    sm.cell(kr, 1, "Key:").font = Font(name=F, sz=9, bold=True)
    for col, k in enumerate(["Pass", "Fail", "In progress", "N/A"], 2):
        bg, fc = STATUS_COLOURS[k]
        c = sm.cell(kr, col, k)
        c.font = Font(name=F, sz=9, bold=True, color=fc)
        c.fill = fill(bg)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX

    hr = kr + 2
    for col, h in enumerate(["#", "Category", "Event", "Variant / label", "Trigger", "Params", "Status"], 1):
        c = sm.cell(hr, col, h)
        c.font = Font(name=F, sz=10, bold=True, color=WHITE)
        c.fill = fill(SLATE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    sm.row_dimensions[hr].height = 24

    first_row, cached = hr + 1, {}
    for i, (ev, (scell, _hdr, pf, pl)) in enumerate(zip(events, blocks)):
        rr = first_row + i
        band = PALE if i % 2 == 1 else WHITE
        variant = (ev.get("variant") or "").strip() or NO_VARIANT
        vals = [i + 1, ev["category"], ev["event"], variant, ev.get("trigger") or "",
                f"=COUNTA('{CHECK_SHEET}'!$A${pf}:$A${pl})", f"='{CHECK_SHEET}'!{scell}"]
        for col, v in enumerate(vals, 1):
            # Params and Status are the two formulas; everything else is text.
            c = sm.cell(rr, col, v) if col in (6, 7) else put(sm, rr, col, v)
            c.fill = fill(band)
            c.border = BOX
            if col == 3:
                c.font = Font(name=F, sz=9, bold=True)
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col == 5:
                c.font = Font(name=F, sz=9)
                c.alignment = Alignment(vertical="top", wrap_text=True)
            elif col == 7:
                c.font = Font(name=F, sz=9, bold=True)
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.font = Font(name=F, sz=9)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cached[f"F{rr}"] = sum(1 for p in ev["params"] if p["name"])
        cached[f"G{rr}"] = statuses[scell]
        need = max(lines_for(ev.get("trigger"), SUMMARY_WIDTHS["E"], 9),
                   lines_for(ev["event"], SUMMARY_WIDTHS["C"], 9),
                   lines_for(variant, SUMMARY_WIDTHS["D"], 9))
        sm.row_dimensions[rr].height = height_for(need, 9, 27.75)
    last = first_row + len(events) - 1
    status_rules(sm, f"G{first_row}:G{last}", 9, True, STATUS_COLOURS)

    tr = last + 2
    c = sm.cell(tr, 6, "Passed:")
    c.font = Font(name=F, sz=10, bold=True)
    c.alignment = Alignment(horizontal="right")
    c = sm.cell(tr, 7, f'=COUNTIF(G{first_row}:G{last},"Pass")&" / "&COUNTA(C{first_row}:C{last})')
    c.font = Font(name=F, sz=10, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    passed = sum(1 for _, (scell, *_rest) in zip(events, blocks) if statuses[scell] == "Pass")
    cached[f"G{tr}"] = f"{passed} / {len(events)}"
    sm.sheet_view.showGridLines = False
    return sm, cached


def build(spec):
    """(workbook, cached results, blocks) for a spec. `blocks` is one
    (status_cell, header_row, first_param_row, last_param_row) per event on
    the QA Checklist tab, in spec order. Raises ValueError on a bad spec."""
    problems = check_spec(spec)
    if problems:
        raise ValueError("Spec problems:\n  " + "\n  ".join(problems))
    events = [dict(ev, params=[norm_param(p) for p in ev["params"]]) for ev in spec["events"]]
    wb = openpyxl.Workbook()
    _ws, blocks, check_cached = build_checklist(wb, spec, events)
    _sm, summary_cached = build_summary(wb, spec, events, blocks, check_cached)
    del wb["Sheet"]
    return wb, {CHECK_SHEET: check_cached, SUMMARY_SHEET: summary_cached}, blocks


# ------------------------------------------------------------------ saving
# Linear on purpose: formula text is XML-escaped, so it never holds a "<", and
# "[^<]*" cannot run past its own cell.
_FORMULA_CELL = re.compile(r'<c r="([A-Z]{1,3}\d+)"([^>]*)><f>([^<]*)</f><v(?:\s*/>|></v>)</c>')


def _inject(xml, values):
    def sub(m):
        coord, attrs, formula = m.group(1), m.group(2), m.group(3)
        if coord not in values:
            return m.group(0)
        value = values[coord]
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        if isinstance(value, bool):
            return f'<c r="{coord}"{attrs} t="b"><f>{formula}</f><v>{int(value)}</v></c>'
        if isinstance(value, (int, float)):
            return f'<c r="{coord}"{attrs}><f>{formula}</f><v>{value}</v></c>'
        return f'<c r="{coord}"{attrs} t="str"><f>{formula}</f><v>{_xml_escape(str(value))}</v></c>'
    return _FORMULA_CELL.sub(sub, xml)


def _sheet_parts(archive):
    """{sheet title: part name} from the workbook and its relationships."""
    workbook = archive.read("xl/workbook.xml").decode("utf-8")
    rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {}
    for rel in re.finditer(r"<Relationship\b[^>]*>", rels):
        rid = re.search(r'Id="([^"]+)"', rel.group(0))
        target = re.search(r'Target="([^"]+)"', rel.group(0))
        if rid and target:
            path = target.group(1).lstrip("/")
            targets[rid.group(1)] = path if path.startswith("xl/") else "xl/" + path
    parts = {}
    for sheet in re.finditer(r"<sheet\b[^>]*>", workbook):
        name = re.search(r'name="([^"]*)"', sheet.group(0))
        rid = re.search(r'r:id="([^"]+)"', sheet.group(0))
        if name and rid and rid.group(1) in targets:
            title = (name.group(1).replace("&quot;", '"').replace("&apos;", "'")
                     .replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&"))
            parts[title] = targets[rid.group(1)]
    return parts


def save(wb, cached, target=None):
    """Save `wb` with each formula's result cached. `target` is a path or a
    binary file object; with none, the .xlsx bytes are returned."""
    buf = io.BytesIO()
    wb.save(buf)
    source = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    parts = _sheet_parts(source)
    by_part = {parts[t]: v for t, v in cached.items() if t in parts and v}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in by_part:
                data = _inject(data.decode("utf-8"), by_part[item.filename]).encode("utf-8")
            zout.writestr(item, data)
    payload = out.getvalue()
    if target is None:
        return payload
    if hasattr(target, "write"):
        target.write(payload)
    else:
        with open(target, "wb") as fh:
            fh.write(payload)
    return None


def write(spec, target=None):
    """Build and save in one step."""
    wb, cached, _blocks = build(spec)
    return save(wb, cached, target)
