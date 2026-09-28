#!/usr/bin/env python3
"""Read a dataLayer QA checklist workbook back into a spec JSON.

Usage:  python checklist_to_spec.py <checklist.xlsx> <spec.json>

Reads any workbook in the QA checklist layout: the one Tapa's DataLayer Doc
Creator exports, the QA team's own template (filled in by hand or built from a
guide), or a previous run of this skill. The result is the expected-parameter
spec for a QA run: every event, its category, variant and trigger, and each
parameter's type, required flag and expected value. Checks and notes already in
the file come across too, under present / value_ok / notes, so a partly tested
checklist is not lost.

Summary rows and QA Checklist blocks are paired by position. The parameter
rows of each block are read from its status formula's range, so rows a tester
inserted inside a block are included.
"""
import json
import re
import sys

import openpyxl

TITLE_SPLIT = re.compile(r"\s+[·—–-]\s+dataLayer QA", re.I)
BLOCK_RANGE = re.compile(r"COUNTA\(\$?A\$?(\d+):\$?A\$?(\d+)\)")
NOTE_SPLIT = "   |   Note: "


def _s(value):
    return "" if value is None else str(value).strip()


def _client(title):
    title = _s(title)
    parts = TITLE_SPLIT.split(title, maxsplit=1)
    return parts[0].strip() if len(parts) > 1 else title


def _batch_and_source(strap):
    batch, source = "", ""
    for part in [p.strip() for p in _s(strap).split("·")]:
        if part.lower().startswith("source:"):
            source = part[len("source:"):].strip()
        elif part.lower().startswith("batch"):
            batch = part
    return batch, source


def _blocks(ws):
    """(header_row, first, last) for each block, top to bottom."""
    found = []
    for row in range(1, ws.max_row + 1):
        formula = ws.cell(row, 6).value
        if isinstance(formula, str) and formula.startswith("=IF(COUNTA("):
            m = BLOCK_RANGE.search(formula)
            if m:
                found.append((row, int(m.group(1)), int(m.group(2))))
    return found


def _summary_rows(sm):
    """[(category, event, variant, trigger)] in table order: the rows under
    the '#' header whose Status cell points at the QA Checklist."""
    rows, header = [], None
    for row in range(1, sm.max_row + 1):
        if _s(sm.cell(row, 1).value) == "#" and _s(sm.cell(row, 3).value) == "Event":
            header = row
            break
    if header is None:
        return rows
    for row in range(header + 1, sm.max_row + 1):
        status = sm.cell(row, 7).value
        if not (isinstance(status, str) and "QA Checklist" in status):
            break
        rows.append((_s(sm.cell(row, 2).value), _s(sm.cell(row, 3).value),
                     _s(sm.cell(row, 4).value), _s(sm.cell(row, 5).value)))
    return rows


def read(path):
    wb = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    ws, sm = wb["QA Checklist"], wb["Summary"]
    client = _client(sm.cell(1, 1).value) or "Client"
    batch, source = _batch_and_source(ws.cell(2, 1).value)
    summary = _summary_rows(sm)

    events = []
    for i, (header, first, last) in enumerate(_blocks(ws)):
        category, event, variant, trigger = summary[i] if i < len(summary) else ("", "", "", "")
        if not event:
            # A block with no Summary row: take the title ("3.  add_to_cart   ·   Hold").
            title = _s(values["QA Checklist"].cell(header, 1).value)
            title = re.sub(r"^\d+\.\s+", "", title)
            event, _, variant = [x.strip() for x in title.partition("   ·   ")]
            category = _s(ws.cell(header, 5).value)
        strip = ws.cell(header + 1, 1).value
        if not isinstance(strip, str) or strip.startswith("="):
            strip = _s(values["QA Checklist"].cell(header + 1, 1).value)
        note = strip.split(NOTE_SPLIT, 1)[1].strip() if NOTE_SPLIT in strip else ""
        if not trigger:
            trigger = strip.split(NOTE_SPLIT, 1)[0].replace("Trigger:", "", 1).strip()

        params = []
        for row in range(first, last + 1):
            name = _s(ws.cell(row, 1).value)
            if not name:
                continue
            params.append({
                "name": name,
                "type": _s(ws.cell(row, 2).value) or "String",
                "required": _s(ws.cell(row, 3).value) or "Yes",
                "expected": _s(ws.cell(row, 4).value),
                "present": _s(ws.cell(row, 5).value),
                "value_ok": _s(ws.cell(row, 6).value),
                "notes": _s(ws.cell(row, 7).value),
            })
        if not event or not params:
            continue
        events.append({"category": category, "event": event,
                       "variant": "" if variant in ("-", "—") else variant,
                       "trigger": trigger, "note": note, "params": params})
    return {"client": client, "batch": batch, "source": source, "events": events}


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    spec = read(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as fh:
        json.dump(spec, fh, indent=1, ensure_ascii=False)
    n = sum(len(e["params"]) for e in spec["events"])
    print(f"{sys.argv[2]}: {spec['client']}, {len(spec['events'])} events, {n} params")


if __name__ == "__main__":
    main()
